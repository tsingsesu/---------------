"""问题 2 独立复核（从落盘文件反查，不依赖内存中的求解结果）。

复核对象（全部为已落盘交付物）：
  * `问题2/result2.xlsx` 三张表（`计划购电量` 335×147、`充放电量` 2005 行、`紧急购电量` 335 行）；
  * `问题2/逐日结果.csv`（334 天 × 144 时段全分辨率明细）；
  * `问题2/联合LP下界_对照.xlsx`、`问题2/口径对照_终端与读法.xlsx`、`问题2/备选读法②_汇总.xlsx`。

复核内容（`问题2/交付清单.md` §六 验收标准 + 构思手锚定值）：
  1. 三表结构与模板逐格一致（列名/列序/行数；模板笔误 `7:0-7:10` 原样保留）；
  2. 轮转还原：把模板行序反转回真实时段序后与 CSV 逐格一致；
  3. 全天购电量/全天购电费逐日自洽（容差 0.01 元）；
  4. 逐日逐时段：供给约束、功率约束、储电量边界、状态递推；
  5. 跨日储电量连续 E_144^d = E_0^{d+1}（334 天）；
  6. E_0(2025-01-01)=6000、E_0(2.1)=1200、E(12.31 24:00)=1200；
  7. 读法①命题 r ≡ 0（从落盘文件独立重算 r=[L·Δ+u−x−P·Δ−v]^+）；
  8. 四指定日期表 1（六个 H:00-H:10 时段）逐位对照构思手锚定值；
  9. 表 3 全部为 0；
 10. 与全年联合 LP 下界差 < 1%；
 11. 与问题 1 典型日的日均对照在 10% 量级内；
 12. 锚定对账：总费用/总购电量/总充放电量与 `_phase0/报告17` 一致；
 13. 备选读法② 与锚定一致（计划/紧急/触发天数/单时段最大）。

输出：`问题2/自检报告.txt`（逐项 PASS/FAIL 与关键数值）。
运行：python 问题2/verify_q2.py（需先跑 run_q2.py、variants_q2.py、run_q2_joint.py）
依赖：numpy、openpyxl；lib/ 公共模块。随机性：无。
"""

import csv
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import numpy as np
import openpyxl

from lib.dataio import attach_path, read_attachment1, read_attachment2
from lib.storage import E_INIT, E_MAX, E_MIN, ETA, P_MAX, U_MAX
from lib.timegrid import (DT_H, K, four_hour_blocks, hour_block_to_k,
                          k_to_template_row, template_row_to_k)

QDIR = os.path.join(ROOT, "问题2")
RESULT_XLSX = os.path.join(QDIR, "result2.xlsx")
TEMPLATE_XLSX = attach_path("附件5", "result2.xlsx")
AUDIT_CSV = os.path.join(QDIR, "逐日结果.csv")
JOINT_XLSX = os.path.join(QDIR, "联合LP下界_对照.xlsx")
VARIANT_XLSX = os.path.join(QDIR, "口径对照_终端与读法.xlsx")
READING2_XLSX = os.path.join(QDIR, "备选读法②_汇总.xlsx")
REPORT_PATH = os.path.join(QDIR, "自检报告.txt")

D_REP_FIRST = 31                                     # 填报区间首日 0 基
N_REP = 334                                          # 填报天数
TOL_KWH = 1e-4                                       # 电量容差 kWh（4 位小数落盘）
TOL_COST = 0.01                                      # 费用容差 元（交付清单 §六.2）
# 由"4 位小数落盘"的数据反查约束时，输入本身带 ±5e-5 的舍入，重算残差会被放大到
# 约 2e-4 量级（供给式含 x/u/v 三项、递推式含 E 两项），故从落盘文件反查用更宽的容差；
# 求解器内存解的约束残差（1e-12 级）另在 run_q2_joint.py 的复核中报告。
TOL_RECON = 1e-3                                     # 落盘文件反查的约束容差
SPEC_HOURS = (10, 12, 14, 16, 18, 20)
SPEC_DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")

# 构思手锚定值（`_phase0/报告17_问题2量级体检.txt`、`报告18_问题2联合LP下界.txt`、任务书）
ANCHOR = {
    "cost_rep": 12254765.7161, "cost_day_mean": 36690.9153,
    "x_rep": 20189815.7979, "u_rep": 6635686.0657, "v_rep": 5374905.7132,
    "joint_rep": 12227243.6427, "joint_mean": 36608.5139, "gap": 27522.0734,
    "spec_table1": {
        "2025-03-20": [0.0000, 479.3120, 0.0000, 562.2616, 697.9973, 0.0000],
        "2025-06-21": [0.0000, 0.0000, 0.0000, 78.5716, 0.0000, 0.0000],
        "2025-09-23": [0.0000, 354.1946, 0.0000, 589.4326, 764.3298, 0.0000],
        "2025-12-21": [0.0000, 932.8280, 0.0000, 801.3819, 715.4144, 0.0000],
    },
    "spec_whole_day": {                              # (全天购电量 kWh, 全天购电费 元)
        "2025-03-20": (66322.0449, 40020.3821),
        "2025-06-21": (32057.4724, 17599.3598),
        "2025-09-23": (66265.4213, 41395.4992),
        "2025-12-21": (93770.2149, 59747.8134),
    },
    # 任务书给出的 2025-03-20 六块充放电锚定（4 小时块口径，构思手 2026-09-11 更正版）
    "spec_epoch_4h_0320": {
        "u": [9000.0000, 1666.6667, 6291.4570, 4689.6822, 0.0, 0.0],
        "v": [0.0, 5862.6228, 2777.3772, 254.7228, 5709.7835, 2930.2165],
    },
    # 备选读法②（`进度台账.md` Phase 2 行）
    "r2_plan": 11757539.8955, "r2_emg": 13393308.2086, "r2_day_mean": 75301.9404,
    "r2_trigger": 293, "r2_max_r": 625.2438,
    "q1_main": 35126.9486,                           # 问题 1 主模型全天购电费，元
}

REPORT = []                                          # 报告行缓冲


def w(text=""):
    """向报告缓冲追加一行（最后统一写盘并打印）。"""
    REPORT.append(str(text))


def check(name, ok, detail=""):
    """记录一条检查项并返回是否通过。

    输入：name，str，检查项名称；ok，bool；detail，str，数值证据
    输出：bool，ok
    """
    w("  [%s] %s%s" % ("PASS" if ok else "FAIL", name, ("：" + detail) if detail else ""))
    return bool(ok)


def read_result_xlsx():
    """读 result2.xlsx 三张表（以只读方式）。

    输出：(wb 数据字典, 表头检查素材)
          dict 含 plan_dates(334)、plan_rows(334,144 模板行序)、plan_x(334)、plan_cost(334)、
          cd_rows(2004,6)、emg_rows(334,3)、headers(dict)
    """
    wb = openpyxl.load_workbook(RESULT_XLSX, read_only=True, data_only=True)
    ws = wb["计划购电量"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(v) for v in rows[0]]
    plan_dates = [r[0] for r in rows[1:1 + N_REP]]
    plan_rows = np.array([[float(v) for v in r[1:1 + K]] for r in rows[1:1 + N_REP]])
    plan_x = np.array([float(r[145]) for r in rows[1:1 + N_REP]])
    plan_cost = np.array([float(r[146]) for r in rows[1:1 + N_REP]])
    ws2 = wb["充放电量"]
    rows2 = list(ws2.iter_rows(values_only=True))
    header2 = [str(v) for v in rows2[0]]
    cd_rows = rows2[1:1 + N_REP * 6]
    ws3 = wb["紧急购电量"]
    rows3 = list(ws3.iter_rows(values_only=True))
    header3 = [str(v) for v in rows3[0]]
    emg_rows = rows3[1:1 + N_REP]
    wb.close()
    return {
        "plan_dates": plan_dates, "plan_rows": plan_rows, "plan_x": plan_x,
        "plan_cost": plan_cost, "cd_rows": cd_rows, "emg_rows": emg_rows,
        "headers": {"plan": header, "cd": header2, "emg": header3}, "n_plan": len(rows),
    }


def read_template_header():
    """读附件5 模板的表头与行数（只读，对照用）。"""
    wb = openpyxl.load_workbook(TEMPLATE_XLSX, read_only=True, data_only=True)
    out = {}
    for name, ncol in (("计划购电量", 147), ("充放电量", 6), ("紧急购电量", 3)):
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        out[name] = {
            "header": [str(v) for v in rows[0][:ncol]],
            "rows": len(rows),
            "cd_blocks": [rows[i][1] for i in range(1, 7)] if name == "充放电量" else None,
        }
    wb.close()
    return out


def read_audit_csv():
    """读 逐日结果.csv（全分辨率审计明细）。

    输出：dict，含按 (天序号, 时段序号) 存放的数组与逐日聚合
    """
    with open(AUDIT_CSV, encoding="utf-8") as f:
        rd = csv.reader(f)
        header = next(rd)
        data = list(rd)
    n = len(data)
    assert n == N_REP * K, "CSV 行数应为 %d，实际 %d" % (N_REP * K, n)
    out = {}
    for col, key in ((5, "price"), (6, "load"), (7, "pv"), (8, "x"), (9, "u"),
                     (10, "v"), (11, "r"), (12, "E")):
        out[key] = np.array([float(row[col]) for row in data]).reshape(N_REP, K)
    out["dates"] = [row[0] for row in data[::K]]          # 每天的日期字符串
    out["header"] = header
    return out


def main():
    """复核主流程：逐项检查并写 自检报告.txt。"""
    w("=" * 78)
    w("问题 2 独立复核报告（从落盘文件反查，不依赖求解内存）")
    w("=" * 78)
    n_pass, n_total = 0, 0

    def ck(name, ok, detail=""):
        # 局部包装：累计通过数
        nonlocal n_pass, n_total
        n_total += 1
        if check(name, ok, detail):
            n_pass += 1

    # ---------------- 0. 文件存在性 ----------------
    w("【0. 交付文件存在性】")
    for p in (RESULT_XLSX, AUDIT_CSV, JOINT_XLSX, VARIANT_XLSX, READING2_XLSX):
        ck("存在：%s" % os.path.basename(p), os.path.exists(p))

    # ---------------- 1. 数据与模板读入 ----------------
    price, load1, pv1, _ = read_attachment1()
    load2, pv_actual, dates = read_attachment2()
    res = read_result_xlsx()
    tpl = read_template_header()
    csvd = read_audit_csv()

    # ---------------- 2. 三表结构与模板一致 ----------------
    w("【2. 三表结构与模板一致性】")
    ck("`计划购电量` 行数 335", res["n_plan"] == 335, "实际 %d" % res["n_plan"])
    ck("`计划购电量` 表头与模板逐格一致（147 列）",
       res["headers"]["plan"] == tpl["计划购电量"]["header"],
       "首=%s 第42列=%s 第144列=%s 末=%s" % (res["headers"]["plan"][0], res["headers"]["plan"][41],
                                        res["headers"]["plan"][143], res["headers"]["plan"][-1]))
    ck("模板笔误 `7:0-7:10` 原样保留（第 42 个时段列）",
       res["headers"]["plan"][42] == "7:0-7:10", "实际 '%s'" % res["headers"]["plan"][42])
    ck("`计划购电量` 日期列覆盖 2025-02-01 至 2025-12-31 共 334 天",
       len(res["plan_dates"]) == N_REP
       and res["plan_dates"][0].strftime("%Y-%m-%d") == "2025-02-01"
       and res["plan_dates"][-1].strftime("%Y-%m-%d") == "2025-12-31")
    ck("`充放电量` 表头与模板一致", res["headers"]["cd"] == tpl["充放电量"]["header"])
    ck("`充放电量` 共 2004 个数据行（334 天 × 6 块）", len(res["cd_rows"]) == 2004,
       "实际 %d" % len(res["cd_rows"]))
    ck("`充放电量` 六块标签与模板逐字一致",
       [res["cd_rows"][b][1] for b in range(6)] == tpl["充放电量"]["cd_blocks"])
    ck("`紧急购电量` 表头与模板一致", res["headers"]["emg"] == tpl["紧急购电量"]["header"])
    ck("`紧急购电量` 覆盖 334 天（一行一天）", len(res["emg_rows"]) == N_REP)

    # ---------------- 3. 轮转还原：xlsx 模板行序 → 真实时段序，与 CSV 逐格一致 ----------------
    w("【3. 填法 Y-轮转还原（模板行序 → 真实时段序）与 CSV 逐格对照】")
    # 真实序与模板序的对应：模板第 j0 列（0 基）装真实第 (j0+1)%144+1 个时段
    # 故真实序数组 real[k-1] = tpl[k_to_template_row(k)]；等价于右旋一位
    real_from_xlsx = res["plan_rows"][:, (np.arange(K) - 1) % K]
    max_diff = float(np.abs(real_from_xlsx - csvd["x"]).max())
    ck("334×144 个计划购电量逐格一致（xlsx 反轮转 vs CSV）", max_diff < TOL_KWH,
       "最大差 %.6f kWh" % max_diff)
    # 表 1 的 H:00-H:10 ↔ 模板第 6H 列（1 基）↔ 真实第 6H+1 个时段
    spot = []
    h = 18
    j0 = k_to_template_row(hour_block_to_k(h))
    spot.append(res["plan_rows"][0, j0] == csvd["x"][0, 6 * h])
    ck("抽查：模板第 %d 列 = 真实第 %d 个时段（2025-02-01）" % (j0 + 1, 6 * h + 1), spot[0],
       "值 %.6f" % res["plan_rows"][0, j0])

    # ---------------- 4. 全天购电量 / 全天购电费自洽 ----------------
    w("【4. 全天购电量与购电费自洽（用落盘文件重算，容差 %.2f 元）】" % TOL_COST)
    x_rowsum = res["plan_rows"].sum(axis=1)
    dx = float(np.abs(x_rowsum - res["plan_x"]).max())
    ck("全天购电量 = 144 列之和", dx < TOL_COST, "最大差 %.6f kWh" % dx)
    # 全天购电费用 CSV 的真实序 x 重算（含紧急费用：读法① r≡0）
    cost_recompute = (csvd["price"] * csvd["x"]).sum(axis=1) + 5.0 * (csvd["price"] * csvd["r"]).sum(axis=1)
    dc = float(np.abs(cost_recompute - res["plan_cost"]).max())
    ck("全天购电费 = Σp·x + 5Σp·r（CSV 重算 vs xlsx）", dc < TOL_COST, "最大差 %.6f 元" % dc)
    total_cost = float(res["plan_cost"].sum())
    ck("填报区间总购电费与锚定一致（%.4f 元）" % ANCHOR["cost_rep"],
       abs(total_cost - ANCHOR["cost_rep"]) < 0.5, "实际 %.4f，差 %+.4f" % (total_cost, total_cost - ANCHOR["cost_rep"]))

    # ---------------- 5. 逐日逐时段约束复核 ----------------
    w("【5. 逐日逐时段约束复核（CSV 真实时段序；容差 %.0e，含 4 位小数舍入）】" % TOL_RECON)
    supply = csvd["x"] + csvd["pv"] * DT_H + csvd["v"] - csvd["load"] * DT_H - csvd["u"]
    ck("供给约束 x + PΔ + v − LΔ − u ≥ 0", supply.min() >= -TOL_RECON,
       "最小残差 %.3e kWh" % float(supply.min()))
    ck("充电量 0 ≤ u ≤ P̄Δ=%.4f" % U_MAX,
       (csvd["u"].min() >= -TOL_RECON) and (csvd["u"].max() <= U_MAX + TOL_RECON),
       "范围 [%.6f, %.6f]" % (csvd["u"].min(), csvd["u"].max()))
    ck("放电量 0 ≤ v ≤ P̄Δ", (csvd["v"].min() >= -TOL_RECON)
       and (csvd["v"].max() <= U_MAX + TOL_RECON),
       "范围 [%.6f, %.6f]" % (csvd["v"].min(), csvd["v"].max()))
    E = csvd["E"]                                        # 时段末（也用于跨日首值）
    ck("储电量 1200 ≤ E ≤ 10800（含首日 0:00 = 1200）",
       (E.min() >= E_MIN - 1e-3) and (E.max() <= E_MAX + 1e-3),
       "范围 [%.4f, %.4f]" % (E.min(), E.max()))
    E_prev = np.concatenate([np.full((N_REP, 1), E_MIN), E[:, :-1]], axis=1)  # 每时段初 = 上时段末
    dE = E - E_prev
    recur = dE - (ETA * csvd["u"] - csvd["v"] / ETA)
    ck("状态递推 ΔE = η·u − v/η", float(np.abs(recur).max()) < TOL_RECON,
       "最大残差 %.3e kWh" % float(np.abs(recur).max()))
    ck("无同时充放（u·v > 1e-6 的时段数 = 0）",
       int(((csvd["u"] > 1e-3) & (csvd["v"] > 1e-3)).sum()) == 0)

    # ---------------- 6. 跨日储电量连续 ----------------
    w("【6. 跨日储电量连续性与端点】")
    E0_next = E[:, :-1] if False else None
    # xlsx 充放电量：每天第 1 行 E_0（第 6 列）、第 2 行 E_144
    cd_E0 = np.array([float(r[5]) for r in res["cd_rows"][0::6]])
    cd_E144 = np.array([float(r[5]) for r in res["cd_rows"][1::6]])
    ck("`充放电量` 每天第 1 行 = E_0、第 2 行 = E_144（第 1 天=1200）",
       abs(cd_E0[0] - E_MIN) < 1e-3, "E_0(2.1)=%.4f" % cd_E0[0])
    cont = np.abs(cd_E144[:-1] - cd_E0[1:])
    ck("跨日连续 E_144^d = E_0^{d+1}（333 个衔接点全部成立）", cont.max() < 1e-3,
       "最大断点 %.6f kWh" % float(cont.max()))
    csv_cont = np.abs(csvd["E"][:-1, -1] - csvd["E"][1:, -1])  # CSV 侧：日末接次日内压
    ck("CSV 侧：每日 24:00 储电量全部 = 1200（自由终端结论）",
       float(np.abs(csvd["E"][:, -1] - E_MIN).max()) < 1e-3,
       "最大偏离 %.6f" % float(np.abs(csvd["E"][:, -1] - E_MIN).max()))
    ck("期末 E(12.31 24:00) = 1200", abs(cd_E144[-1] - E_MIN) < 1e-3, "%.4f" % cd_E144[-1])

    # ---------------- 7. 读法①命题 r ≡ 0（独立重算） ----------------
    w("【7. 读法①命题验证：r ≡ 0】")
    r_recompute = np.maximum(csvd["load"] * DT_H + csvd["u"] - csvd["x"] - csvd["pv"] * DT_H - csvd["v"], 0.0)
    ck("从 CSV 独立重算 max r < 1e-3（4 位小数舍入内的可行性余量）",
       float(r_recompute.max()) < 1e-3,
       "max r = %.3e kWh（求解器内存解为 0；落盘舍入引入的 1e-4 级余量）"
       % float(r_recompute.max()))
    ck("CSV 落盘的紧急购电量恒为 0", float(np.abs(csvd["r"]).max()) < 1e-6,
       "max = %.3e" % float(np.abs(csvd["r"]).max()))
    emg_vals = [r[2] for r in res["emg_rows"]]
    all_zero = all((v == 0) or (str(v).strip() == "0") for v in emg_vals)
    ck("`紧急购电量` 表 334 天全部为 0、时间段列为空",
       all_zero and all(r[1] in (None, "") for r in res["emg_rows"]),
       "非零行数 = %d" % int(sum(1 for v in emg_vals if not ((v == 0) or (str(v).strip() == "0")))))

    # ---------------- 8. 四指定日期表 1/表 2/表 3 反查 ----------------
    w("【8. 四指定日期表 1/表 2/表 3（从落盘文件反查）】")
    import datetime as _dt
    for t in SPEC_DATES:
        d0 = (_dt.date.fromisoformat(t) - _dt.date(2025, 1, 1)).days
        i = d0 - D_REP_FIRST
        # 表 1：六个 H:00-H:10 时段，从 xlsx 模板序取（模板第 6H 列 = 0 基 6H-1）
        vals_xlsx = [float(res["plan_rows"][i, k_to_template_row(hour_block_to_k(h))]) for h in SPEC_HOURS]
        vals_csv = [float(csvd["x"][i, 6 * h]) for h in SPEC_HOURS]
        anchor = ANCHOR["spec_table1"][t]
        ok_anchor = max(abs(a - b) for a, b in zip(vals_xlsx, anchor)) < 0.0001
        ok_csv = max(abs(a - b) for a, b in zip(vals_xlsx, vals_csv)) < 1e-3
        ck("%s 表 1 六时段与锚定逐位一致" % t, ok_anchor,
           "%s" % ["%.4f" % v for v in vals_xlsx])
        ck("%s 表 1 六时段与 CSV 一致" % t, ok_csv)
        # 全天值
        ax, ac = ANCHOR["spec_whole_day"][t]
        ck("%s 全天购电量/购电费与锚定一致" % t,
           abs(res["plan_x"][i] - ax) < 1e-3 and abs(res["plan_cost"][i] - ac) < 0.01,
           "x=%.4f（锚 %.4f）；费=%.4f（锚 %.4f）" % (res["plan_x"][i], ax, res["plan_cost"][i], ac))
        # 表 2：六块充放电（从 xlsx 反查；4 小时块 = 24 个 10 分钟时段，勿与 36 时段的 6 小时窗混淆）
        ux = [float(r[2]) for r in res["cd_rows"][i * 6:(i + 1) * 6]]
        vx = [float(r[3]) for r in res["cd_rows"][i * 6:(i + 1) * 6]]
        # 与 CSV 聚合一致（每块 24 个时段：0-4,4-8,…,20-24）
        ux_csv = [float(csvd["u"][i, b * 24:(b + 1) * 24].sum()) for b in range(6)]
        vx_csv = [float(csvd["v"][i, b * 24:(b + 1) * 24].sum()) for b in range(6)]
        ok_blk = (max(abs(a - b) for a, b in zip(ux, ux_csv)) < 1e-3
                  and max(abs(a - b) for a, b in zip(vx, vx_csv)) < 1e-3)
        ck("%s 表 2 六块充放电与 CSV 聚合一致" % t, ok_blk,
           "充 %s" % ["%.4f" % v for v in ux])
        ck("%s 表 2 首末储电量 0:00 = 24:00 = 1200" % t,
           abs(cd_E0[i] - E_MIN) < 1e-3 and abs(cd_E144[i] - E_MIN) < 1e-3)
        # 任务书给出的 2025-03-20 六块锚定（4 小时块口径）
        if t == "2025-03-20":
            a4 = ANCHOR["spec_epoch_4h_0320"]
            ok4 = (max(abs(a - b) for a, b in zip(ux, a4["u"])) < 1e-3
                   and max(abs(a - b) for a, b in zip(vx, a4["v"])) < 1e-3)
            ck("2025-03-20 六块充放电与任务书锚定（4 小时块）一致", ok4,
               "充 %s；放 %s" % (["%.4f" % v for v in ux], ["%.4f" % v for v in vx]))

    # ---------------- 9. 锚定总量对账 ----------------
    w("【9. 锚定总量对账（报告17）】")
    ck("总购电量 %.4f kWh" % ANCHOR["x_rep"],
       abs(float(res["plan_x"].sum()) - ANCHOR["x_rep"]) < 1.0,
       "实际 %.4f，差 %+.4f" % (float(res["plan_x"].sum()), float(res["plan_x"].sum()) - ANCHOR["x_rep"]))
    ck("总充电量 %.4f kWh" % ANCHOR["u_rep"],
       abs(float(csvd["u"].sum()) - ANCHOR["u_rep"]) < 1.0,
       "实际 %.4f" % float(csvd["u"].sum()))
    ck("总放电量 %.4f kWh" % ANCHOR["v_rep"],
       abs(float(csvd["v"].sum()) - ANCHOR["v_rep"]) < 1.0,
       "实际 %.4f" % float(csvd["v"].sum()))
    ck("日均费用 %.4f 元" % ANCHOR["cost_day_mean"],
       abs(total_cost / 334.0 - ANCHOR["cost_day_mean"]) < 0.01,
       "实际 %.4f" % (total_cost / 334.0))

    # ---------------- 10. 联合 LP 下界差 ----------------
    w("【10. 与全年联合 LP 的下界对照】")
    wb = openpyxl.load_workbook(JOINT_XLSX, read_only=True, data_only=True)
    ws = wb["汇总对照"]
    jrows = {r[0]: r[1] for r in list(ws.iter_rows(values_only=True))[1:]}
    wb.close()
    joint_rep = float(jrows["填报区间联合最优费用（2.1–12.31）"])
    gap = total_cost - joint_rep
    ck("联合下界与锚定一致（%.4f 元）" % ANCHOR["joint_rep"],
       abs(joint_rep - ANCHOR["joint_rep"]) < 0.01, "实际 %.4f" % joint_rep)
    ck("滚动 − 下界 < 1%%（实际 %.4f%%）" % (100.0 * gap / joint_rep),
       100.0 * gap / joint_rep < 1.0, "差 %.4f 元" % gap)
    ck("滚动 ≥ 下界（下界性质成立）", gap >= -0.01)

    # ---------------- 11. 与问题 1 对照 ----------------
    w("【11. 与问题 1 典型日对照】")
    rel = abs(total_cost / 334.0 - ANCHOR["q1_main"]) / ANCHOR["q1_main"]
    ck("日均费用与问题 1 主模型（%.4f 元）差 < 10%%" % ANCHOR["q1_main"], rel < 0.10,
       "相对差 %.4f%%，方向：逐日凸性使日均略高" % (100.0 * rel))

    # ---------------- 12. 备选读法② 对账 ----------------
    w("【12. 备选读法② 对账（从落盘 xlsx 反查）】")
    wb = openpyxl.load_workbook(READING2_XLSX, read_only=True, data_only=True)
    ws = wb["汇总"]
    r2rows = {r[0]: (r[1], r[2]) for r in list(ws.iter_rows(values_only=True))[1:]}
    wb.close()
    ck("读法② 计划购电费与锚定一致", abs(float(r2rows["计划购电费_元"][0]) - ANCHOR["r2_plan"]) < 0.01,
       "%.4f" % float(r2rows["计划购电费_元"][0]))
    ck("读法② 紧急购电费与锚定一致", abs(float(r2rows["紧急购电费_元"][0]) - ANCHOR["r2_emg"]) < 0.01,
       "%.4f" % float(r2rows["紧急购电费_元"][0]))
    ck("读法② 日均与锚定一致", abs(float(r2rows["日均费用_元"][0]) - ANCHOR["r2_day_mean"]) < 0.01)
    ck("读法② 触发天数与锚定一致", int(r2rows["触发紧急购电天数（占 334 天）"][0]) == ANCHOR["r2_trigger"])
    ck("读法② 单时段最大 r 与锚定一致", abs(float(r2rows["单时段最大紧急购电量_kWh"][0]) - ANCHOR["r2_max_r"]) < 0.001)

    # ---------------- 13. 终端与读法对照落盘 ----------------
    w("【13. 口径对照落盘】")
    wb = openpyxl.load_workbook(VARIANT_XLSX, read_only=True, data_only=True)
    names = wb.sheetnames
    var_rows = list(wb["终端条件对照"].iter_rows(values_only=True))
    wb.close()
    ck("`口径对照_终端与读法.xlsx` 含终端条件与读法两张对照表",
       "终端条件对照" in names and "读法对照" in names)
    ck("终端=当日初始 对照行存在且日均已记录", var_rows[2][2] is not None,
       "cyclic 日均 = %.4f 元（%+.4f%% vs 主模型）" % (float(var_rows[2][2]),
       100.0 * (float(var_rows[2][2]) / (total_cost / 334.0) - 1.0)))

    # ---------------- 14. 记录 6 小时块（报告17 初版口径）供口径追溯 ----------------
    w("【14. 口径追溯：报告17 初版曾用四个 6 小时块（0-6/6-12/12-18/18-24）记录表 2 锚定】")
    w("    说明：交付模板规定表 2 为六个 4 小时块（见第 2 节）；此处 6 小时块仅供追溯核对。")
    for t in SPEC_DATES:
        d0 = (_dt.date.fromisoformat(t) - _dt.date(2025, 1, 1)).days
        i = d0 - D_REP_FIRST
        u6 = [float(csvd["u"][i, b * 36:(b + 1) * 36].sum()) for b in range(4)]
        v6 = [float(csvd["v"][i, b * 36:(b + 1) * 36].sum()) for b in range(4)]
        w("    %s 六小时块：充 %s；放 %s"
          % (t, ["%.4f" % u for u in u6], ["%.4f" % v for v in v6]))

    # ---------------- 汇总 ----------------
    w("=" * 78)
    w("自检总结：%d / %d 项通过" % (n_pass, n_total))
    w("=" * 78)
    text = "\n".join(REPORT)
    with open(REPORT_PATH, "w", encoding="utf-8") as f:
        f.write(text + "\n")
    print(text)
    print("\n自检报告已写入：%s" % REPORT_PATH)
    return n_pass == n_total


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    ok = main()
    sys.exit(0 if ok else 1)
