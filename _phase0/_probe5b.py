import openpyxl, os
W = r"C:\Users\lhl87\Desktop\高教社杯全国大学生数学建模大赛"
lines = []
for fn in ["result1.xlsx","result2.xlsx","result3.xlsx","result4-2.xlsx","result4-3.xlsx"]:
    wb = openpyxl.load_workbook(os.path.join(W,"附件","附件5",fn))
    for ws in wb.worksheets:
        fmts = {}
        for row in ws.iter_rows(min_row=1, max_row=min(3,ws.max_row)):
            for c in row:
                fmts[c.number_format] = fmts.get(c.number_format,0)+1
        lines.append("%s [%s] rows=%d cols=%d 数字格式=%s" % (fn, ws.title, ws.max_row, ws.max_column, fmts))
    wb.close()
open(os.path.join(W,"_phase0","报告06_格式.txt"),"w",encoding="utf-8").write("\n".join(lines))

# 抽取 result2 计划购电量表头标签，单独成文件
wb = openpyxl.load_workbook(os.path.join(W,"附件","附件5","result2.xlsx"))
ws = wb["计划购电量"]
hdr = [c.value for c in ws[1]]
out = ["标签总数=%d" % len(hdr)]
for i,h in enumerate(hdr,1):
    out.append("[%3d] %r" % (i,h))
# 充放电量 / 紧急购电量 的完整内容
ws2 = wb["充放电量"]
out.append("--- 充放电量 ---")
for r in ws2.iter_rows(values_only=True):
    out.append(repr(r))
ws3 = wb["紧急购电量"]
out.append("--- 紧急购电量 ---")
for r in ws3.iter_rows(values_only=True):
    out.append(repr(r))
open(os.path.join(W,"_phase0","报告07_result2明细.txt"),"w",encoding="utf-8").write("\n".join(out))

wb = openpyxl.load_workbook(os.path.join(W,"附件","附件5","result1.xlsx"))
out = []
for nm in ["计划购电量","充放电量"]:
    ws = wb[nm]
    out.append("--- %s ---" % nm)
    for r in ws.iter_rows(values_only=True):
        out.append(repr(r))
open(os.path.join(W,"_phase0","报告08_result1明细.txt"),"w",encoding="utf-8").write("\n".join(out))
print("ok")
