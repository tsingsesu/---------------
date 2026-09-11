"""Phase 0 探针 5：结果模板的逐行逐列精确抄录（供交付清单引用）。"""

import os
import openpyxl

WORK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = lambda *a: os.path.join(WORK, *a)
out = []
def w(s=""):
    out.append(str(s))


for fn in ["result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"]:
    w("=" * 78)
    w("文件: 附件/附件5/%s" % fn)
    wb = openpyxl.load_workbook(P("附件", "附件5", fn), read_only=False, data_only=False)
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        w("  工作表 [%s]: 行=%d 列=%d" % (ws.title, len(rows), ws.max_column))
        w("    数字格式样例: A1=%r B1=%r B2=%r" % (ws["A1"].number_format, ws["B1"].number_format,
                                                 ws["B2"].number_format if ws.max_row > 1 else None))
        if ws.title in ("计划购电量", "调整购电量"):
            hdr = rows[0]
            w("    时间列标签(共%d个):" % (len(hdr) - 1))
            for i, h in enumerate(hdr[1:], start=1):
                w("       [%3d] %r" % (i, h))
            w("    日期列: 共%d行数据, 首=%r 末=%r" % (len(rows) - 1, rows[1][0], rows[-1][0]))
        else:
            for i, r in enumerate(rows):
                w("     [%2d] %s" % (i, r))
    wb.close()

with open(P("_phase0", "报告05_模板逐条.txt"), "w", encoding="utf-8") as f:
    f.write("\n".join(out))
print("done")
