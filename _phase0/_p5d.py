import openpyxl, os
W = r"C:\Users\lhl87\Desktop\高教社杯全国大学生数学建模大赛"
lines=[]
for fn in ["result2.xlsx","result3.xlsx"]:
    wb = openpyxl.load_workbook(os.path.join(W,"附件","附件5",fn))
    ws = wb["紧急购电量"]
    lines.append("=== %s 紧急购电量 rows=%d cols=%d" % (fn, ws.max_row, ws.max_column))
    for r in ws.iter_rows(values_only=True):
        lines.append("   " + repr(r))
    ws = wb["充放电量"]
    lines.append("=== %s 充放电量 rows=%d cols=%d 末3行" % (fn, ws.max_row, ws.max_column))
    rows=list(ws.iter_rows(values_only=True))
    for r in rows[-3:]:
        lines.append("   " + repr(r))
    ws = wb["调整购电量"] if "调整购电量" in wb.sheetnames else None
    if ws is not None:
        lines.append("=== %s 调整购电量 rows=%d cols=%d" % (fn, ws.max_row, ws.max_column))
        lines.append("    hdr前3=%s 后4=%s" % ([c.value for c in ws[1]][:3], [c.value for c in ws[1]][-4:]))
        lines.append("    日期列 首=%r 末=%r" % (ws.cell(2,1).value, ws.cell(ws.max_row,1).value))
    wb.close()
open(os.path.join(W,"_phase0","报告10_紧急购电样本.txt"),"w",encoding="utf-8").write("\n".join(lines))
print("ok")
