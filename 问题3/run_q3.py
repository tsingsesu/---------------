r"""问题 3 主脚本：多阶段滚动 LP（0:00 计划 + 6:00/12:00/18:00 调整），逐日 4 个串联 LP。

流程：读附件1 电价 + 附件2 负载/光伏实际 + 附件3 整点预报 → 按 D-07 主口径 M2 分解到
      144 个 10 分钟值 → 从 2025-01-01（E_0=6000 kWh）逐日多阶段滚动 365 天（1 月预热）
      → 落盘 `全分辨率明细.csv` 与 `result3.xlsx` 四张表 → 打印关键数值与自检结果。

口径（全部来自 `口径与假设台账.md`，不得自行更改）：
  * D-01 填法 Y-轮转：模板第 i 格装当天第 (i+1)%144+1 个时段；第 144 格在时间上最早；
  * D-04 终端储电量自由 + 跨日传递 E_{144}^d = E_0^{d+1}，E_0^1 = 6000 kWh；
  * D-05 单向效率 η=0.9；
  * D-07 预报分解主口径 M2（整点线性插值，lib/forecast.py）；
  * D-11 从 2025-01-01 仿真、1 月预热，result3.xlsx 只填 2025-02-01 至 12-31 共 334 天；
  * D-12 结算 J = Σ[p·min(x,y) + κ₋p(x−y)⁺ + κ₊p(y−x)⁺] + κΣp·r，κ₋=0.5、κ₊=1.5、κ=5；
  * D-13 调整不可追溯：k∈(0:00,6:00] 用计划量；k∈(6:00,12:00] 用 6:00 决策值；余类推；
  * D-14 `调整购电量` 填最终生效总量 y（不是增量）；
  * D-15 统一 4 位小数；D-16.8/9/10 模板扩表、紧急购电写法、`7:0-7:10` 笔误照抄。

运行：python 问题3/run_q3.py
依赖：numpy、scipy、openpyxl（均本机已装）；lib/ 公共模块（solve_day/run_days/forecast/xlsxio）。
随机性：无（确定性 LP 求解，不需要随机种子）。
"""

import os
import shutil
import sys

# 把工作区根目录加入模块搜索路径，保证在任意工作目录下都能 from lib... import ...
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import numpy as np
import openpyxl

from lib.dataio import attach_path, read_attachment1, read_attachment2
from lib.forecast import load_forecast
from lib.logio import Tee
from lib.run_days import solve_rolling_staged
from lib.solve_day import KAPPA_EMG, KAPPA_OVER, KAPPA_UNDER, settle_total
from lib.storage import E_INIT, E_MAX, E_MIN, ETA, P_MAX
from lib.timegrid import (DT_H, K, four_hour_blocks, hour_block_to_k,
                          hour_block_to_template_row, k_to_label)
from lib.xlsxio import (emergency_text, fill_row_periods, r4, set_decimal_format,
                        template_order, wb_date)

QDIR = os.path.join(ROOT, "问题3")                                   # 问题 3 交付目录
TEMPLATE_XLSX = attach_path("附件5", "result3.xlsx")                 # 题目模板（只读）
RESULT_XLSX = os.path.join(QDIR, "result3.xlsx")                     # 交付结果文件
AUDIT_CSV = os.path.join(QDIR, "全分辨率明细.csv")                     # 全分辨率审计明细（334 天）
LOG_PATH = os.path.join(QDIR, "主模型运行日志.txt")                   # 运行日志（数值证据）

ND = 4                                                               # 统一小数位数（D-15）
D_REP_FIRST = 31                                                     # 填报区间首日：2025-02-01（0 基 31）
N_REP = 334                                                          # 填报区间天数（2025-02-01 至 12-31）
SPEC_HOURS = (10, 12, 14, 16, 18, 20)                                # 题目表 1 的六个指定时段
SPEC_DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")  # 表 1/2/3 的四个指定日期
STAGES_MAIN = (0, 6, 12, 18)                                         # 主模型：全调整

# 构思手独立锚定值（`_phase0/报告19_问题3量级体检.txt`，2026-09-11 版；探针偏差行有符号
# 笔误，已将修正后的重跑值一并列出，详见 `编程手汇报.md` 的假设说明——仅作对照打印）
ANCHOR = {
    "deg": 12254765.7161,      # 退化自检：预报=实际 时应复现的问题 2 总费用，元
    "a": 14741164.4870,        # 策略 (a) 仅 0:00（探针值，与修正值一致到 0.3 元）
    "b": 14211401.0661,        # 策略 (b) +6:00（探针偏差行修正后重跑值），元
    "c": 13816123.4566,        # 策略 (c) +6/12（修正后重跑值），元
    "d": 13816127.9140,        # 策略 (d) 全调整（修正后重跑值），元
}

# 全分辨率明细 CSV 的列定义（自解释、供外部审计逐条核验）
CSV_HEADER = ["日期", "天序号", "时段序号", "时段起", "时段止", "电价_元每kWh",
              "负载_kW", "光伏实际_kW", "光伏预报0_kW", "光伏预报6_kW", "光伏预报12_kW",
              "光伏预报18_kW", "计划购电量_kWh", "调整购电量_kWh", "欠取量_kWh", "超用量_kWh",
              "充电量_kWh", "放电量_kWh", "紧急购电量_kWh", "弃光电量_kWh", "时段末储电量_kWh"]


def write_audit_csv(path, dates, price, load, pv_actual, pv_fc144, roll):
    """落盘全分辨率审计明细 CSV：填报区间 334 天 × 144 时段 × 21 列。

    输入：path，str，目标 csv 路径；dates，list[date]，365 天日期
          price (K,) 元/kWh；load/pv_actual (D,K) kW；pv_fc144 (D,4,K) kW
          roll，dict，solve_rolling_staged 的返回值（填报区间部分）
    输出：str，写入的路径
    """
    lines = [",".join(CSV_HEADER)]
    for i, d in enumerate(roll["d_index"]):
        for k in range(1, K + 1):
            start_label, end_label = k_to_label(k)          # 该时段的（起, 止）钟点
            row = [
                dates[d].isoformat(), str(d + 1), str(k), start_label, end_label,
                "%.4f" % price[k - 1],
                "%.4f" % load[d, k - 1], "%.4f" % pv_actual[d, k - 1],
                "%.4f" % pv_fc144[d, 0, k - 1], "%.4f" % pv_fc144[d, 1, k - 1],
                "%.4f" % pv_fc144[d, 2, k - 1], "%.4f" % pv_fc144[d, 3, k - 1],
                "%.4f" % r4(roll["x_plan"][i, k - 1]),
                "%.4f" % r4(roll["y_adj"][i, k - 1]),
                "%.4f" % r4(roll["d_under"][i, k - 1]),
                "%.4f" % r4(roll["d_over"][i, k - 1]),
                "%.4f" % r4(roll["u_chg"][i, k - 1]),
                "%.4f" % r4(roll["v_dis"][i, k - 1]),
                "%.4f" % r4(roll["r_emg"][i, k - 1]),
                "%.4f" % r4(roll["g_curt"][i, k - 1]),
                "%.4f" % r4(roll["E_soc"][i, k]),
            ]
            lines.append(",".join(row))
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write("\n".join(lines) + "\n")
    return path


def fill_result_xlsx(path, dates, roll, price):
    """把多阶段滚动解按模板规格填入 result3.xlsx（先复制模板，绝不改写附件）。

    四张表：
      `计划购电量`：0:00 计划 x（335×147，轮转填法）；`全天购电量`=Σx、`全天购电费`=J_plan=Σp·x；
      `调整购电量`：最终生效总量 y（D-14；同构），`全天购电量`=Σy、`全天购电费`=J（含调整与
      紧急的全天总费用，D-12 口径；两张表的列标签相同，口径差异写进问题3/问题3.md 表注）；
      `充放电量`：最终充放电量与每日 0:00/24:00 储电量（334 天 × 6 个 4 小时块）；
      `紧急购电量`：逐日紧急购电的时间段与电量文本（334 行）。

    输入：path，str，目标 result3.xlsx；dates，list[date]；roll，dict（填报区间）；price (K,)
    输出：dict，写入摘要（四个指定日期的表 1/2/3 值与紧急购电统计）
    """
    shutil.copyfile(TEMPLATE_XLSX, path)                # 附件目录只读（D-16.6）
    wb = openpyxl.load_workbook(path)                   # 载入副本，保留模板全部标签

    # ---------------- 工作表 1：计划购电量 ----------------
    ws = wb["计划购电量"]
    for i, d in enumerate(roll["d_index"]):
        row = 2 + i
        assert ws.cell(row=row, column=1).value.date() == dates[d], "模板日期列与填报区间不一致"
        fill_row_periods(ws, row, roll["x_plan"][i])          # 轮转填法（D-01）+ 4 位小数
        c_x = ws.cell(row=row, column=146)
        c_x.value = float(r4(roll["x_plan"][i].sum()))        # 全天购电量，kWh
        set_decimal_format(c_x)
        c_j = ws.cell(row=row, column=147)
        c_j.value = float(r4(roll["J_plan"][i]))              # 全天购电费 = 计划购电费用 Σp·x，元
        set_decimal_format(c_j)

    # ---------------- 工作表 2：调整购电量（与计划表同构，填最终生效总量） ----------------
    ws2 = wb["调整购电量"]
    for i, d in enumerate(roll["d_index"]):
        row = 2 + i
        fill_row_periods(ws2, row, roll["y_adj"][i])
        c_x = ws2.cell(row=row, column=146)
        c_x.value = float(r4(roll["y_adj"][i].sum()))         # 全天购电量（调整后），kWh
        set_decimal_format(c_x)
        c_j = ws2.cell(row=row, column=147)
        c_j.value = float(r4(roll["J_day"][i]))               # 全天购电费 = 总费用
        set_decimal_format(c_j)                               # （计划+调整+紧急，D-12 口径）

    # ---------------- 工作表 3：充放电量（334 天 × 6 个 4 小时块 = 2004 行） ----------------
    ws3 = wb["充放电量"]
    blocks = four_hour_blocks()                          # 六个 4 小时块（真实时间顺序）
    for i, d in enumerate(roll["d_index"]):
        for b, (start_label, end_label, k_first, k_last) in enumerate(blocks):
            row = 2 + i * 6 + b
            if b == 0:                                   # 日期只填每天第一行（与模板一致）
                dc = ws3.cell(row=row, column=1)
                dc.value = wb_date(dates[d], ws3.cell(row=2, column=1))
                dc.number_format = ws3.cell(row=2, column=1).number_format
            ws3.cell(row=row, column=2).value = ws3.cell(row=2 + b, column=2).value  # 块标签照抄
            cu = ws3.cell(row=row, column=3)
            cu.value = float(r4(np.sum(roll["u_chg"][i, k_first - 1:k_last])))
            set_decimal_format(cu)
            cv = ws3.cell(row=row, column=4)
            cv.value = float(r4(np.sum(roll["v_dis"][i, k_first - 1:k_last])))
            set_decimal_format(cv)
            ct = ws3.cell(row=row, column=5)
            if b == 0:
                ct.value = ws3.cell(row=2, column=5).value          # datetime.time(0, 0)
            elif b == 1:
                ct.value = ws3.cell(row=3, column=5).value          # 字符串 '24:00'
            cs = ws3.cell(row=row, column=6)
            if b == 0:
                cs.value = float(r4(roll["E_soc"][i, 0]))
                set_decimal_format(cs)
            elif b == 1:
                cs.value = float(r4(roll["E_soc"][i, K]))
                set_decimal_format(cs)

    # ---------------- 工作表 4：紧急购电量（334 天，每天一行） ----------------
    ws4 = wb["紧急购电量"]
    n_trigger = 0
    spec_emg = {}
    for i, d in enumerate(roll["d_index"]):
        row = 2 + i
        dc = ws4.cell(row=row, column=1)
        dc.value = wb_date(dates[d], ws4.cell(row=2, column=1))
        dc.number_format = ws4.cell(row=2, column=1).number_format
        time_text, energy_text = emergency_text(roll["r_emg"][i], k_to_label)
        ws4.cell(row=row, column=2).value = time_text
        ws4.cell(row=row, column=3).value = energy_text if time_text != "" else 0
        if time_text != "":
            n_trigger += 1
        if dates[d].isoformat() in SPEC_DATES:
            spec_emg[dates[d].isoformat()] = (time_text, energy_text)

    wb.save(path)                                       # 保存副本，附件模板不受影响
    return {"n_trigger": n_trigger, "spec_emg": spec_emg}


def print_spec_tables(dates, roll, summary):
    """控制台打印四个指定日期的表 1/2/3 数值（供构思手核对与审计追溯）。

    输入：dates，list[date]；roll，dict；summary，fill_result_xlsx 的返回摘要
    输出：无（仅打印）
    """
    for t in SPEC_DATES:
        import datetime as _dt
        d0 = (_dt.date.fromisoformat(t) - _dt.date(2025, 1, 1)).days
        i = d0 - D_REP_FIRST
        print("  --- %s（天序号 %d）---" % (t, d0 + 1))
        for h in SPEC_HOURS:
            k = hour_block_to_k(h)
            j0 = hour_block_to_template_row(h)
            print("    %2d:00-%2d:10  真实时段 %3d  模板列 %3d   y = %10.4f kWh"
                  % (h, h, k, j0 + 1, roll["y_adj"][i, k - 1]))
        print("    全天购电量(y) = %.4f kWh；计划购电费 = %.4f 元；全天总费用 J = %.4f 元"
              % (float(roll["y_adj"][i].sum()), float(roll["J_plan"][i]), float(roll["J_day"][i])))
        print("    调整相关费用 = %.4f 元；紧急购电费 = %.4f 元；Σr = %.4f kWh"
              % (float(roll["J_adj"][i]), float(roll["J_emg"][i]), float(roll["r_emg"][i].sum())))
        for start_label, end_label, k_first, k_last in four_hour_blocks():
            print("    %s-%s  充 = %10.4f  放 = %10.4f"
                  % (start_label, end_label,
                     float(np.sum(roll["u_chg"][i, k_first - 1:k_last])),
                     float(np.sum(roll["v_dis"][i, k_first - 1:k_last]))))
        print("    0:00 储电量 = %.4f kWh；24:00 储电量 = %.4f kWh"
              % (roll["E_soc"][i, 0], roll["E_soc"][i, K]))
        print("    表 3（紧急购电）：时间段='%s'，购电量='%s'"
              % summary["spec_emg"].get(t, ("", "0")))


def main():
    """问题 3 主流程：载入数据 → 预报分解 → 全年多阶段滚动 → 落盘 → 打印关键数值。"""
    print("=" * 78)
    print("问题 3：多阶段滚动 LP（0:00 计划 + 6:00/12:00/18:00 调整；D-07 M2 分解；D-12 结算）")
    print("=" * 78)

    # ---------------- 1. 读数据 ----------------
    price, load1, pv1, _ = read_attachment1()            # 附件1：问题 3 只用其电价列（D-08）
    load, pv_actual, dates = read_attachment2()          # 附件2：负载（视为已知）与光伏实际（仅评估）
    pv_fc, pv_fc144, dates3, tau_list = load_forecast(method="linear")
    assert load.shape == (365, K) and pv_actual.shape == (365, K)
    assert pv_fc144.shape == (365, 4, K)
    print("附件1 电价：%.4f–%.4f 元/kWh；附件2：%d 天 × %d 时段（%s 至 %s）"
          % (price.min(), price.max(), 365, K, dates[0], dates[-1]))
    print("附件3 预报：%d 天 × %d 个发布时刻 × 24 整点 → 分解（M2 线性插值）为 %d 个 10 分钟值"
          % (pv_fc.shape[0], pv_fc.shape[1], K))

    # ---------------- 2. 全年多阶段滚动（含 1 月预热） ----------------
    roll_all = solve_rolling_staged(price, load, pv_actual, pv_fc144,
                                    stages=STAGES_MAIN, d_start=0, d_end=365)
    sl = slice(D_REP_FIRST, 365)
    keys = ("d_index", "x_plan", "y_adj", "u_chg", "v_dis", "r_emg", "g_curt", "E_soc",
            "d_under", "d_over", "J_plan", "J_adj", "J_emg", "J_day", "J_alt_day")
    roll = {key: roll_all[key][sl] for key in keys}
    print("多阶段滚动完成：365 天全部阶段 status 之和为 0 = %s"
          % bool(np.all(roll_all["status_day"] == 0)))

    # ---------------- 3. 落盘 ----------------
    write_audit_csv(AUDIT_CSV, dates, price, load, pv_actual, pv_fc144, roll)
    summary = fill_result_xlsx(RESULT_XLSX, dates, roll, price)
    print("已落盘：%s（334 天 × 144 时段 × %d 列）" % (AUDIT_CSV, len(CSV_HEADER)))
    print("已落盘：%s（由附件5 模板复制后填写四张表，附件未改动）" % RESULT_XLSX)

    # ---------------- 4. 关键数值打印 ----------------
    tot_plan = float(roll["J_plan"].sum())
    tot_adj = float(roll["J_adj"].sum())
    tot_emg = float(roll["J_emg"].sum())
    tot = float(roll["J_day"].sum())
    tot_alt = float(roll["J_alt_day"].sum())
    print("-" * 78)
    print("【全年汇总（填报区间 334 天，2025-02-01 至 12-31）】")
    print("计划购电费用 J_plan = %.4f 元" % tot_plan)
    print("调整相关费用 J_adj  = %.4f 元（欠取为负、超用为正）" % tot_adj)
    print("紧急购电费用 J_emg  = %.4f 元" % tot_emg)
    print("总费用 J            = %.4f 元；日均 = %.4f 元" % (tot, tot / N_REP))
    print("对照口径总费用（欠取按全价+违约金）= %.4f 元（差 %+.4f 元）" % (tot_alt, tot_alt - tot))
    print("总计划量 Σx = %.4f kWh；总最终量 Σy = %.4f kWh；调整量 Σ|y−x| = %.4f kWh"
          % (float(roll["x_plan"].sum()), float(roll["y_adj"].sum()),
             float(np.sum(roll["d_under"] + roll["d_over"]))))
    print("紧急购电量 Σr = %.4f kWh；触发紧急购电天数 = %d / %d；单时段最大 r = %.4f kWh"
          % (float(roll["r_emg"].sum()), summary["n_trigger"], N_REP, float(roll["r_emg"].max())))
    print("填报区间期初 E_0(2.1 0:00) = %.4f kWh；期末 E(12.31 24:00) = %.4f kWh"
          % (roll["E_soc"][0, 0], roll["E_soc"][-1, K]))
    print("与问题 2（完全信息）对照：问题 2 = %.4f 元；预报与调整机制的总代价 = %+.4f 元（%+.4f%%）"
          % (ANCHOR["deg"], tot - ANCHOR["deg"], 100.0 * (tot / ANCHOR["deg"] - 1.0)))

    print("-" * 78)
    print("【题目表 1/表 2/表 3：四个指定日期（真实时段序 k=1..144；H:00-H:10 ↔ 第 6H+1 个时段）】")
    print_spec_tables(dates, roll, summary)

    print("-" * 78)
    print("【自检】")
    # 非追溯性：k<=36（0:00–6:00）必须 y == x
    dev_early = float(np.abs(roll["y_adj"][:, :36] - roll["x_plan"][:, :36]).max())
    print("非追溯性：k<=36 的 max|y−x| = %.2e（应为 0）" % dev_early)
    # 储电量边界
    print("储电量轨迹范围：min = %.4f，max = %.4f（应在 [%.0f, %.0f] 内）"
          % (float(roll["E_soc"].min()), float(roll["E_soc"].max()), E_MIN, E_MAX))
    # 跨日连续性：前一日 24:00 == 后一日 0:00
    gap = float(np.abs(roll["E_soc"][:-1, K] - roll["E_soc"][1:, 0]).max())
    print("跨日连续性：max|E_144^d − E_0^{d+1}| = %.2e（应为 0）" % gap)
    # 抽样：用落盘前的数组独立重算一天的总费用（与聚合口径互检）
    i0 = 50                                              # 任取一天做独立重算演示
    st_i = settle_total(price, roll["x_plan"][i0], roll["y_adj"][i0], roll["r_emg"][i0])
    print("独立重算示例（第 %d 个填报日）：J = %.4f（聚合数组为 %.4f，差 %.2e）"
          % (i0 + 1, st_i["J"], roll["J_day"][i0], abs(st_i["J"] - roll["J_day"][i0])))
    print("=" * 78)
    return roll, summary


if __name__ == "__main__":
    tee = Tee(LOG_PATH)
    original_stdout = sys.stdout
    sys.stdout = tee
    try:
        main()
    finally:
        sys.stdout = original_stdout
        tee.close()
    print("运行日志已写入：%s" % LOG_PATH)
